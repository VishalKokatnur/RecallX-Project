"""Extract slide and table text from PowerPoint presentations."""

from pptx import Presentation


def extract_text_from_pptx(file_path: str) -> str:
    presentation = Presentation(file_path)
    slides = []

    for slide in presentation.slides:
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))

        if parts:
            slides.append("\n".join(parts))

    return "\n\n".join(slides)
